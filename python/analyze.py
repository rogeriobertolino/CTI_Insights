"""Análise XLSX com leitura portátil e estatística descritiva da biblioteca padrão."""
import json
import math
import re
import statistics
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def slug(value):
    text = unicodedata.normalize("NFD", str(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def find(columns, aliases):
    normalized = [slug(column) for column in columns]
    for alias in aliases:
        if alias in normalized:
            return normalized.index(alias)
    for alias in aliases:
        for index, column in enumerate(normalized):
            if alias in column:
                return index
    return -1


def distribution(rows, index, split=False):
    if index < 0:
        return []
    values = []
    for row in rows:
        items = re.split(r"[,;|/]", row[index]) if split else [row[index]]
        values.extend(value.strip() for value in items if value.strip())
    return [{"name": name, "value": count} for name, count in Counter(values).most_common(20)]


def normalized_segment(value):
    return {
        "ind": "Indústria", "industria": "Indústria", "industrial": "Indústria",
        "com": "Comércio", "comercio": "Comércio", "serv": "Serviços",
        "servico": "Serviços", "servicos": "Serviços",
    }.get(slug(value), str(value).strip())


def normalized_level(value):
    return str(value).strip().upper()


def parse_money(value):
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if not text:
        return None
    if "," in text and "." in text:
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        group = "." if decimal == "," else ","
        parts = text.split(decimal)
        if len(parts) != 2 or not re.fullmatch(r"\d{1,2}", parts[1]) or not re.fullmatch(
            r"[+-]?\d{1,3}(?:" + re.escape(group) + r"\d{3})+", parts[0]
        ):
            return None
        text = parts[0].replace(group, "") + "." + parts[1]
    elif "," in text:
        if not re.fullmatch(r"[+-]?\d+,\d{1,2}", text):
            return None
        text = text.replace(",", ".")
    if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", text):
        return None
    number = float(text)
    return number if math.isfinite(number) else None


def quantile(values, fraction):
    position = (len(values) - 1) * fraction
    lower, upper = math.floor(position), math.ceil(position)
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def descriptive_statistics(rows, index):
    if index < 0:
        return {}
    values = sorted(value for row in rows if (value := parse_money(row[index])) is not None)
    if not values:
        return {}
    return {
        "count": len(values), "total": round(math.fsum(values), 2),
        "mean": round(statistics.mean(values), 2), "median": round(statistics.median(values), 2),
        "std": round(statistics.stdev(values), 2) if len(values) > 1 else 0,
        "min": round(values[0], 2), "max": round(values[-1], 2),
        "variance": statistics.variance(values) if len(values)>1 else None, "modes": statistics.multimode(values), "range": values[-1]-values[0], "q2": statistics.median(values),
        "q1": round(quantile(values, .25), 2), "q3": round(quantile(values, .75), 2),
    }


def cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        pattern = cell.number_format or ""
        if re.fullmatch(r"0+", pattern) and float(value).is_integer():
            return str(int(value)).zfill(len(pattern))
        if float(value).is_integer():
            return str(int(value))
    return str(value).strip()


def read_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["dados_importacao"] if "dados_importacao" in workbook.sheetnames else workbook.worksheets[0]
        if (sheet.max_row or 0) > 10001 or (sheet.max_column or 0) > 100:
            raise ValueError("A planilha deve ter até 10.000 registros e 100 colunas.")
        matrix = [[cell_text(cell) for cell in row] for row in sheet.iter_rows()]
    finally:
        workbook.close()
    while matrix and not any(matrix[-1]):
        matrix.pop()
    while matrix and not any(matrix[0]):
        matrix.pop(0)
    if not matrix:
        raise ValueError("A planilha está vazia.")
    headers = []
    for index, column in enumerate(matrix[0]):
        base = column or f"Coluna {index + 1}"
        name, suffix = base, 2
        while name in headers:
            name = f"{base} ({suffix})"
            suffix += 1
        headers.append(name)
    rows = [(row + [""] * len(headers))[:len(headers)] for row in matrix[1:]]
    return headers, rows


def analyze(path):
    headers, rows = read_rows(path)
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "analytics-python"))
    from pipeline import executar_etl
    analytics = executar_etl([dict(zip(headers, row)) for row in rows], headers)
    if analytics["schema"] == "comercial-contratos-v1":
        headers = list(analytics["validRecords"][0]) if analytics["validRecords"] else headers
        rows = [[str(record.get(h, "")) for h in headers] for record in analytics["validRecords"]]
    indexes = {
        "clientCode": find(headers, ["cliente_codigo", "codigo_cliente", "cod_cliente", "codigo", "cnpj"]),
        "clientName": find(headers, ["razao_social", "nome_cliente", "nome", "cliente", "empresa"]),
        "segment": find(headers, ["segmento", "setor"]), "status": next((i for i,h in enumerate(headers) if slug(h) in ("status", "situacao")), -1),
        "service": find(headers, ["servico", "produto", "solucao"]),
        "consultant": find(headers, ["consultor", "responsavel", "vendedor"]),
        "level": find(headers, ["nivel", "classificacao"]),
        "revenue": find(headers, ["faturamento_anual", "faturamento_mensal", "faturamento", "receita"]),
    }
    client = indexes["clientCode"] if indexes["clientCode"] >= 0 else indexes["clientName"]
    keys = [slug(row[client]) for row in rows] if client >= 0 else []
    non_empty = [key for key in keys if key]
    duplicates = len(non_empty) - len(set(non_empty))
    missing_key = keys.count("") if client >= 0 else len(rows)
    empty_cells = sum(value == "" for row in rows for value in row)
    completeness = max(0, 100 - empty_cells / max(len(rows) * len(headers), 1) * 100)
    quality = max(0, completeness - duplicates / max(len(rows), 1) * 20 - missing_key / max(len(rows), 1) * 30)
    active = sum(row[indexes["status"]].lower() == "ativo" for row in rows) if indexes["status"] >= 0 else None
    for row in rows:
        if indexes["segment"] >= 0:
            row[indexes["segment"]] = normalized_segment(row[indexes["segment"]])
        if indexes["level"] >= 0:
            row[indexes["level"]] = normalized_level(row[indexes["level"]])
    return {
        "headers": headers, "records": [dict(zip(headers, row)) for row in rows], "totalRows": len(rows),
        "totalClients": len(set(non_empty)) + missing_key if client >= 0 else len(rows), "activeClients": active,
        "completeness": round(completeness, 2), "quality": round(quality, 2),
        "duplicates": duplicates, "missingKey": missing_key, "emptyCells": empty_cells,
        "segments": distribution(rows, indexes["segment"]), "statuses": distribution(rows, indexes["status"]),
        "services": distribution(rows, indexes["service"], True), "consultants": distribution(rows, indexes["consultant"]),
        "levels": distribution(rows, indexes["level"]), "revenueStatistics": descriptive_statistics(rows, indexes["revenue"]),
        "columns": indexes, "analytics": analytics,
    }


def main(path):
    print(json.dumps(analyze(path), ensure_ascii=False, allow_nan=False))


if __name__ == "__main__":
    if len(sys.argv) != 2 or not Path(sys.argv[1]).is_file():
        raise SystemExit("Uso: analyze.py arquivo.xlsx")
    main(sys.argv[1])
