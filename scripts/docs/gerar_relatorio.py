"""Gera o relatório acadêmico da atividade de validação com Pinia.

Dependência: python-docx. Execute com o Python do ambiente de documentos.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[2] / "docs" / "Relatorio_Atividade_Pinia.docx"


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.06
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Title"].paragraph_format.space_after = Pt(3)
    doc.styles["Subtitle"].font.size = Pt(12)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(13)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.core_properties.title = "Validação de planilhas com Pinia"
    doc.core_properties.subject = "Atividade individual de Front End do CTI Provedor"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "Vue 3, Pinia, Excel, validação"


def paragraph(doc, text, bold_start=None):
    p = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold = True
        p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    for index, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for col, cell in enumerate(row.cells):
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{edge}")
                for key, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
                    border.set(qn(f"w:{key}"), value)
                borders.append(border)
            tc_pr.append(borders)
            margin = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:w"), "95")
                el.set(qn("w:type"), "dxa")
                margin.append(el)
            tc_pr.append(margin)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EEF5" if index == 0 else "FFFFFF")
            tc_pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.03
                for r in p.runs:
                    r.font.size = Pt(10.5)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.bold = index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def build():
    doc = Document()
    configure(doc)
    doc.add_paragraph("Validação de planilhas com Pinia", "Title")
    doc.add_paragraph("CTI Provedor   Atividade individual de Front End", "Subtitle")
    paragraph(doc, "Nome: ____________________________________    Data: ____ / ____ / ______")
    paragraph(doc,
        "A funcionalidade analisa uma planilha Excel no navegador e apresenta os problemas "
        "encontrados em um relatório. A store usePlanilhaStore, em src/stores/planilha.js, "
        "centraliza os dados e as regras de validação no Vue 3. O fluxo é Excel → tela de "
        "upload → Pinia → tela de relatório.")

    doc.add_paragraph("State e armazenamento temporário", "Heading 1")
    paragraph(doc,
        "nomeArquivo e nomeAba identificam a análise. cabecalhos guarda os nomes das "
        "colunas reconhecidas. registros contém os valores lidos e a linha original do "
        "Excel, no formato { linha, dados }.")
    paragraph(doc,
        "erros armazena cada ocorrência com linha, campo, tipo, descricao e valor "
        "original, além de sugestao quando aplicável. carregando controla a espera, "
        "erroArquivo informa falhas de leitura ou estrutura e validadoEm registra o "
        "momento da validação.")
    paragraph(doc,
        "O estado fica somente na memória enquanto a aplicação permanece aberta. "
        "Recarregar a página ou usar Limpar remove a análise. O arquivo original é "
        "preservado e esse fluxo não envia os dados ao servidor.")

    doc.add_paragraph("Getters e informações calculadas", "Heading 1")
    table(doc, ["Getter", "Cálculo ou resultado"], [
        ("totalRegistros", "Quantidade de registros analisados, excluindo linhas totalmente vazias."),
        ("totalErros", "Número de ocorrências encontradas em todos os campos."),
        ("totalComErro", "Quantidade de linhas distintas com pelo menos um problema."),
        ("totalValidos", "Total de registros menos o total de linhas com erro."),
        ("percentualValidos", "Percentual de registros válidos; retorna zero se não houver registros."),
        ("errosPorTipo", "Quantidade de cada categoria de erro, inclusive categorias com zero ocorrências."),
        ("registrosValidos", "Lista dos registros que não possuem ocorrências de validação."),
    ], [1.6, 5.3])
    paragraph(doc,
        "Uma linha pode ter vários problemas. Por isso, totalErros pode ser maior que "
        "totalComErro; cada linha é contada apenas uma vez no total de registros com erro.")

    doc.add_page_break()
    doc.add_paragraph("Actions e processamento", "Heading 1")
    paragraph(doc,
        "importarArquivo recebe o arquivo, verifica o formato e o tamanho, lê a "
        "primeira aba com SheetJS, reconhece os cabeçalhos e armazena os registros. "
        "Em seguida, chama validarDados. Arquivos inválidos geram uma mensagem em erroArquivo.",
        "importarArquivo")
    paragraph(doc,
        "validarDados percorre os registros, executa as regras e substitui a lista de "
        "erros. A revalidação não acumula ocorrências antigas nem corrige automaticamente "
        "os valores originais. limpar cancela leituras pendentes e restaura o estado inicial.",
        "validarDados")

    doc.add_paragraph("Validações realizadas", "Heading 1")
    table(doc, ["Validação", "Regra aplicada"], [
        ("Campos obrigatórios", "codigo, nome e email não podem ficar vazios ou conter apenas espaços. status e nivel são opcionais."),
        ("Espaços desnecessários", "Detecta espaços no início e no fim, espaços repetidos, tabulações e quebras de linha em qualquer campo."),
        ("Registros duplicados", "Compara codigo ignorando caixa e espaços extras. Todos os registros do grupo são marcados, inclusive o primeiro."),
        ("E-mails inválidos", "Verifica a estrutura do endereço, o domínio, espaços e pontos consecutivos, além dos limites de tamanho."),
        ("Dados fora do padrão", "Valida os caracteres de codigo e nome. Aceita status Ativo, Inativo ou Pendente e nivel A, B ou C."),
        ("Textos para padronizar", "Sinaliza nomes inteiramente em maiúsculas ou minúsculas, e-mails válidos com maiúsculas e a caixa incorreta de status e nivel. Exibe sugestões."),
    ], [1.6, 5.3])
    paragraph(doc,
        "A importação aceita .xlsx de até 5 MiB, com área utilizada até a linha 10001 "
        "e a coluna 100. Cabeçalhos ausentes, vazios ou duplicados, planilhas sem registros "
        "e arquivos ilegíveis são rejeitados. Linhas totalmente vazias são ignoradas; "
        "linhas contendo espaços são analisadas.")

    doc.add_paragraph("Apresentação dos resultados", "Heading 1")
    paragraph(doc,
        "A tela de upload, em /app/upload, permite selecionar o arquivo, baixar o modelo "
        "e testar um exemplo com problemas. Após a análise, /app/relatorios mostra nome "
        "do arquivo, aba, data e hora, totais de registros, válidos, linhas com erro e ocorrências.")
    paragraph(doc,
        "O relatório reúne as quantidades por tipo e uma tabela pesquisável com linha do "
        "Excel, campo, descrição, valor original e sugestão. A prévia identifica os registros "
        "válidos e os que precisam de revisão. Os controles permitem filtrar, paginar, "
        "revalidar e iniciar outra análise.")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
